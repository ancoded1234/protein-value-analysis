"""Actual Office export and public/local browser checks, with data-linked assertions."""
from pathlib import Path
import os,json
import pandas as pd
import win32com.client
from pptx import Presentation
from PIL import Image
from playwright.sync_api import sync_playwright
R=Path(__file__).resolve().parents[1];Q=R/'analysis/qa';Q.mkdir(exist_ok=True);P=R/'presentation';C=pd.read_csv(R/'data/comparison.csv')
url=os.getenv('PRESENTATION_TEST_URL',(R/'docs/index.html').as_uri());public=url.startswith('https://')
report={'url':url,'desktop':[],'mobile':[],'errors':[]}
if not public:
 app=win32com.client.DispatchEx('PowerPoint.Application');p=app.Presentations.Open(str(P/'protein_value_analysis.pptx'),False,False,False)
 report['office_slides']=p.Slides.Count;assert p.Slides.Count==13;charts=[]
 try:
  for sl in p.Slides:
   sl.SlideShowTransition.EntryEffect=1793
   for sh in sl.Shapes:
    if sh.HasChart:
     charts.append(dict(slide=sl.SlideIndex,values=list(sh.Chart.SeriesCollection(1).Values)))
     fx=sl.TimeLine.MainSequence.AddEffect(sh,22,0,3);fx.Timing.Duration=.9
   sl.Export(str(Q/f'ppt_{sl.SlideIndex:02}.png'),'PNG',1600,900)
  p.Save();p.SaveAs(str(P/'protein_value_analysis.pdf'),32);p.Close()
 finally:app.Quit()
 assert len(charts)==2 and all(abs(a-b)<.001 for a,b in zip(charts[0]['values'],C.cost_20g));report['office_charts']=charts
 app=win32com.client.DispatchEx('PowerPoint.Application');p=app.Presentations.Open(str(P/'protein_value_analysis.pptx'),True,False,False);assert p.Slides.Count==13;p.Close();app.Quit();report['office_reopened']=True
 deck=Presentation(P/'protein_value_analysis.pptx');assert all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)
 sheet=Image.new('RGB',(1200,1125),'#e5eadf')
 for i,f in enumerate(sorted(Q.glob('ppt_*.png'))):
  im=Image.open(f);im.thumbnail((390,219));sheet.paste(im,((i%3)*400,(i//3)*225))
 sheet.save(Q/'contact_sheet.jpg')
with sync_playwright() as pw:
 b=pw.chromium.launch(headless=True);page=b.new_page(viewport={'width':1600,'height':1000});page.on('pageerror',lambda x:report['errors'].append(str(x)));page.goto(url);page.wait_for_timeout(500)
 assert page.locator('.slide').count()==13
 for n in range(13):
  page.evaluate(f'show({n})');page.wait_for_timeout(1300);size=page.locator('.slide.active').evaluate('s=>({x:s.scrollWidth-s.clientWidth,y:s.scrollHeight-s.clientHeight})');report['desktop'].append({'slide':n+1,**size});page.screenshot(path=str(Q/f'web_{n+1:02}.png'))
 page.evaluate('show(0)');page.keyboard.press('ArrowRight');assert page.locator('#page').inner_text()=='02 / 13'
 page.locator('#notes').click();assert page.locator('dialog').evaluate('d=>d.open');page.locator('#close').click();page.locator('#sources').click();assert page.locator('.source').count()>0;page.locator('#close').click()
 page.evaluate('show(6)');page.locator('#budget').fill('1100');page.locator('#budget').dispatch_event('input');assert page.locator('#budgetbars .barrow').count()==1;report['budget_filter']='passed'
 page.evaluate('show(5)');page.locator('#motion').click();page.wait_for_function('document.querySelector("video").readyState>=2');report['video']=page.locator('video').evaluate('v=>({width:v.videoWidth,duration:v.duration,error:v.error?.message??null})');assert not report['video']['error'];page.locator('#close').click()
 page.locator('#fullscreen').click();page.wait_for_function('!!document.fullscreenElement');page.evaluate('document.exitFullscreen()')
 mobile=b.new_page(viewport={'width':390,'height':844},is_mobile=True,has_touch=True);mobile.goto(url)
 for n in range(13):
  mobile.evaluate(f'show({n})');mobile.wait_for_timeout(80);x=mobile.evaluate('document.documentElement.scrollWidth-innerWidth');report['mobile'].append({'slide':n+1,'overflow_x':x});assert x<=1
 mobile.evaluate('show(6)');mobile.wait_for_timeout(1500);mobile.screenshot(path=str(Q/'mobile_value.png'));b.close()
assert not report['errors'];assert all(x['x']<=1 and x['y']<=1 for x in report['desktop']),report['desktop']
report['status']='passed';(Q/('public_browser.json' if public else 'validation.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print('Passed: 13 slides, charts, desktop/mobile, sources, notes, filter, video, fullscreen'+(' + Office open/export/reopen' if not public else ' on public URL'))
