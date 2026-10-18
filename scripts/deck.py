"""One reviewed slide narrative -> editable PPTX and animated HTML presentation."""
from pathlib import Path
import json,html,shutil
import pandas as pd
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION
R=Path(__file__).resolve().parents[1];D=R/'docs';P=R/'presentation';A=D/'assets';A.mkdir(exist_ok=True)
C=pd.read_csv(R/'data/comparison.csv');S=pd.read_csv(R/'data/sources.csv');M=json.loads((R/'analysis/metrics.json').read_text(encoding='utf-8'));L=pd.read_csv(R/'data/lab_2023.csv')
short={'P01':'더단백','P02':'셀렉스','P04':'그린비아','P05':'마이밀','P06':'오트몬드'};C['short']=C.product_id.map(short)
slides=[
 ('PROTEIN / VALUE LAB','단백질 20g, 얼마면 될까?','온라인 판매가격과 영양성분으로 찾는 가성비 프로틴', '초콜릿 계열 RTD 음료 6개 후보 중 5개 구성 비교. 분말과 음료를 섞지 않았다. 오늘의 가격 관측과 2023년 독립 시험을 분리했다.','D01|D02|D04|D05|D06|C01'),
 ('THE ANSWER FIRST','관측한 5개 구성 중, 그린비아가 가장 저렴했다',f"단백질 20g당 약 {M['winner_cost20']:,.0f}원 · 선택한 판매 구성·표시 성분 기준", '그린비아 초코 250mL 18개, 상품가 18,940원·배송비 0원 관측이다. 시장 전체 최저가나 맛·흡수율·안전성 1위라는 뜻은 아니다. 가격 효율과 열량 효율은 다른 기준이다.','D04'),
 ('DEFINE THE SAMPLE','프로틴이라는 이름보다 비교 조건을 먼저 맞춘다','초콜릿 계열 · 바로 마시는 RTD · 영양성분과 판매 구성을 연결', '후보는 더단백, 셀렉스, 닥터유, 그린비아, 마이밀, 오트몬드다. 닥터유는 24g과 27g 제품 세대 연결이 불명확해 순위에서 제외했다. 포함 5개 제품도 묶음 수가 6–24개로 달라 결과는 이 관측 구성에 한정된다.','D01|D02|D03|D04|D05|D06|M02'),
 ('COLLECT & VERIFY','실제 판매 화면, 가격비교 원문, 공식 자료를 대조했다',f"판매처 가격 관측 {M['offers']}건 · 후보 6개 · 계산에 포함 5개", '쿠팡은 computer-use로 로그인 없는 제품 화면을 확인했다. 다나와는 Firecrawl 본문에서 판매처 가격과 배송비를 추출했다. 네이버페이 판매처 값은 다나와 경유이며 네이버 직접 결제 검증은 아니다. 셀렉스 공식 영양표와 한국소비자원 과거 시험을 별도 확보했다. k-skill 쇼핑 모듈을 검토했으나 자동 수집 제한과 API 경로 특성 때문에 호출하지 않았다.','C01|D01|D02|D04|D05|D06|M01|K01'),
 ('THE DENOMINATOR','총액을 단백질 총량으로 나누면 공정해진다','단백질 20g당 비용 = (상품가 + 배송비) ÷ (팩 수 × 팩당 단백질) × 20', '20g은 비교를 위한 공통 분모이며 개인 섭취 권장량이 아니다. 상품가와 배송비를 더하고, 선택한 묶음의 전체 단백질량으로 나눈다. 쿠폰·멤버십 적립금은 현금 가격에서 빼지 않았고 실결제 가격이라고 표현하지 않았다.','D01|D02|D04|D05|D06|C01'),
 ('THE SHIPPING TRAP','21,580원보다 23,900원이 더 저렴할 수 있다','더단백 초코 250mL 18개 · 같은 구성의 배송비 포함 비교', '다나와에 재게시된 11번가 상품가는 21,580원이지만 배송비 3,000원을 더하면 24,580원이다. 쿠팡 직접 화면은 23,900원·무료배송이었다. 따라서 상품가만 고르면 이 사례에서는 680원을 더 지불하게 된다. 두 값 모두 관측 표시가이며 주문 검증은 아니다.','D01|C01'),
 ('COST PER 20G','같은 단백질량으로 맞추면 비용은 최대 2.37배','선택한 판매 구성 중 배송비까지 확인한 최저 관측값', '그린비아 약 1052원, 더단백 1328원, 오트몬드 1768원, 마이밀 2367원, 셀렉스 2491원이다. 이는 각 구성에서 확보한 판매처 중 배송비 포함 최소값이다. 가격비교 화면에서 모든 판매처와 모든 묶음을 수집한 것은 아니다.','D01|D02|D04|D05|D06|C01'),
 ('NUTRITION EFFICIENCY','열량 대비 단백질은 셀렉스가 가장 높았다','가격 효율과 영양성분의 효율을 각각 보여준다', '셀렉스 20g/99kcal는 100kcal당 약 20.2g이다. 그린비아는 20g/120kcal로 약 16.7g이다. 이 비율은 표시 영양성분으로 계산한 지표이며 생체 흡수율이나 전체 품질의 점수가 아니다. 셀렉스 수치는 제조사 공식 자료로도 확인했다.','D01|D02|D04|D05|D06|M01'),
 ('WHAT QUALITY MEANS','상세 영양표까지 확인해야 누락을 줄일 수 있다','더단백 당류 0.8g 반영 · 상품 요약만 보고 미확보로 처리했던 오류 정정', '최초 조사에서 더단백 상품 요약에 당류가 없다는 이유로 상세 이미지까지 확인하지 않고 미확보로 처리했다. 사용자 제공 쿠팡 상세 라벨에서 250mL·105kcal·단백질20g·당류0.8g을 확인해 정정했다. 셀렉스는 공식 0g, 그린비아는 재게시 1g, 오트몬드는 재게시 0g이다. 마이밀은 현재 판매 구성의 상세 라벨 검토가 아직 완료되지 않은 상태이며 정보가 공개되지 않았다는 뜻이 아니다. 100% WPI 원료 사용이라는 말은 음료 전체가 단백질이라는 뜻이 아니며 원료 유형을 우열 점수로 바꾸지 않았다.','C02|D02|D04|D05|D06|M01'),
 ('A TRADE-OFF, NOT A MAGIC SCORE','최저 비용과 최소 열량은 같은 선택이 아니다','왼쪽은 저렴함 · 아래쪽은 같은 단백질량에서 열량이 낮음', '비용과 단백질 20g당 열량을 나란히 비교한다. 가중치를 임의로 정한 종합 품질 점수는 만들지 않았다. 원료 선호, 실제 맛, 소화 경험, 알레르기는 이 차트로 판단할 수 없다.','D01|D02|D04|D05|D06|C01|M01'),
 ('INDEPENDENT TEST / 2023','과거 시험에서도 가격이 곧 품질 지표는 아니었다','한국소비자원 16개 제품 시험 · 2023년 자료, 현재 순위와 분리', '2023년 8월 발표 보고서의 16개 시험 제품 중 14개는 당시 비교 기준인 아미노산스코어 85 이상이었다. 단백질 1g당 보고 가격은 32–375원이었다. 가장 저렴한 NS포대유청 WPC는 아미노산스코어 127이었다. 서로 다른 유형과 용도의 제품이므로 이 사실을 현재 모든 프로틴에 일반화하지 않는다. 당시 가격은 2023년 2월 기준이고 현재 판매 제품에 이 점수를 붙이지 않았다.','K01'),
 ('BUY BY YOUR PRIORITY','가성비는 ‘내가 무엇을 비교하느냐’에서 갈린다','가격은 20g당 비용, 성분은 별도 지표로 판단한다', '이번 관측 구성에서 가격 우선이면 그린비아, 열량 대비 단백질 우선이면 셀렉스가 앞섰다. 이것은 개인별 섭취나 의료 권고가 아닌 비교 데이터의 요약이다. 구매 시에는 같은 맛·용량·단백질 표시·수량·배송비를 다시 확인해야 한다. 결론은 특정 브랜드의 절대적 품질 우위가 아니라 단백질량과 배송비를 맞춰 비교해야 한다는 것이다.','D01|D02|D04|D05|D06|C01|M01'),
 ('SOURCES & LIMITS','가격은 스냅샷, 영양표는 실험 결과와 다르다','수집 2026.10.18 · 재게시 가격은 판매처·실결제와 차이 날 수 있음', '소규모 목적 표본으로 시장 전체를 대표하지 않는다. 6팩과 24팩의 구매 부담과 배송비가 다르다. 표시 성분은 실제 함량 실험이 아니다. 2023년 시험은 과거 제품과 당시 가격에만 적용한다. 전체 CSV·분석 코드·제외 사유를 저장소에 보존하고 원문 전문과 인증 파일은 공개하지 않는다.','|'.join(S.source_id)),
]
NAVY='122523';BG='F3F5EB';GREEN='527858';LIME='C8E85B';MUTED='738177';WHITE='FFFFFF';LINE='D7DED2'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
def color(c):return RGBColor.from_string(c)
def box(sl,x,y,w,h,c):
 sh=sl.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=color(c);sh.line.fill.background();return sh
def txt(sl,t,x,y,w,h=.5,size=18,bold=False,c=NAVY):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=sh.text_frame;f.clear();f.word_wrap=True;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
 for i,line in enumerate(str(t).split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line;p.font.name='맑은 고딕';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c);p.space_after=Pt(8)
 return sh
def chart(sl,data,field,x=.6,y=3.0,w=11.9,h=3.5):
 cd=CategoryChartData();cd.categories=data['short'].tolist();cd.add_series('값',data[field].tolist())
 ch=sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(x),Inches(y),Inches(w),Inches(h),cd).chart;ch.has_legend=False;ch.chart_style=10;ch.has_title=False;ch.category_axis.reverse_order=True;ch.category_axis.tick_labels.font.name='맑은 고딕';ch.category_axis.tick_labels.font.size=Pt(13);ch.value_axis.minimum_scale=0;ch.value_axis.tick_labels.font.size=Pt(10)
 ch.value_axis.major_gridlines.format.line.color.rgb=color(LINE)
 ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;ch.plots[0].data_labels.number_format='#,##0.0' if field=='protein_per_100kcal' else '#,##0';ch.plots[0].data_labels.font.size=Pt(13)
 for i,p in enumerate(ch.series[0].points):p.format.fill.solid();p.format.fill.fore_color.rgb=color(GREEN if i==0 else LIME);p.format.line.fill.background()
 return ch
def card(sl,x,y,w,label,headline,body):
 box(sl,x,y,w,2.75,WHITE);txt(sl,label,x+.22,y+.2,w-.44,.35,10,True,GREEN);txt(sl,headline,x+.22,y+.72,w-.44,.8,27,True);txt(sl,body,x+.22,y+1.7,w-.44,.9,15,False,MUTED)
for n,(tag,title,sub,notes,ids) in enumerate(slides,1):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(NAVY if n in [1,2,12] else BG);dark=n in [1,2,12];fg=BG if dark else NAVY
 txt(sl,tag,.55,.3,11,.3,10,True,LIME if dark else GREEN);txt(sl,title,.55,1.05,12.1,1.14,31,True,fg);txt(sl,sub,.58,2.23,12,.58,13,False,'B8C5B6' if dark else MUTED)
 txt(sl,f'{n:02} / 13',12.05,7.03,.75,.3,10,True,fg);txt(sl,' · '.join(ids.split('|'))+' | 수집 2026.10.18 · 공개 표시값·관측 구성 한정',.55,7.03,11.3,.28,8,False,'B8C5B6' if dark else MUTED)
 refs=S[S.source_id.isin(ids.split('|'))];sl.notes_slide.notes_text_frame.text=notes+'\n\n'+'\n'.join(f'[{r.source_id}] {r.title}\n{r.url}\n게시 {r.published} / 수집 {r.collected}' for r in refs.itertuples())
 if n==1:
  txt(sl,'20g',.6,3.15,7,1.8,106,True,LIME);txt(sl,'같은 단백질량으로 비교하는\n온라인 프로틴 가격 실험',.67,5.45,7.4,.9,22,True,BG)
  for i,h in enumerate([2.5,3.05,2.1]):box(sl,9+i*.83,6.2-h,.67,h,LIME if i==0 else GREEN)
 elif n==2:
  txt(sl,f"{M['winner_cost20']:,.0f}원",.65,3.2,8,1.4,83,True,LIME);txt(sl,'그린비아 초코 250mL × 18팩\n18,940원 ÷ (18 × 20g) × 20g',.7,5.0,8,.9,21,True,BG);txt(sl,'가격 효율 1위\n≠\n전체 품질 1위',9.35,3.6,3.2,2.2,25,True,BG)
 elif n==3:
  for i,r in enumerate(C.itertuples()):
   y=3.05+i*.51;txt(sl,r.short,.7,y,1.8,.35,15,True);txt(sl,f'{r.protein_g}g / {r.volume_ml}mL',2.75,y,3,.35,15);txt(sl,f'{r.pack_count}팩 관측',6,y,2,.35,15);txt(sl,r.nutrition_status,9,y,3,.35,13,False,MUTED)
  txt(sl,'제외: 닥터유 24g·27g 세대 연결 미확정 / 포함 제품도 전국·전체 묶음 대표 표본 아님',.7,6.15,11.8,.55,12,False,MUTED)
 elif n==4:
  for i,(a,b,c) in enumerate([('01 / OBSERVE','쿠팡 실제 화면','로그아웃 가격·배송비·수량'),('02 / EXTRACT','Firecrawl 원문','판매처별 가격·배송비 추출'),('03 / CHECK','공식·공공 자료','표시 성분과 과거 시험 분리')]):card(sl,.6+i*4.2,3.3,3.98,a,b,c)
 elif n==5:
  txt(sl,'상품가 + 배송비',.75,3.25,11.8,.8,38,True);box(sl,.75,4.22,11.65,.025,GREEN);txt(sl,'팩 수 × 팩당 단백질(g)',.75,4.62,9.4,.8,34,True);txt(sl,'× 20',10.2,4.62,2.2,.8,34,True,GREEN);txt(sl,'20g은 비교 단위이며 개인 섭취 권장량이 아니다.',.8,6.05,11.5,.45,15,False,MUTED)
 elif n==6:sl.shapes.add_picture(str(R/'analysis/figures/05_shipping.png'),Inches(.7),Inches(3.0),width=Inches(11.7),height=Inches(3.7))
 elif n==7:chart(sl,C,'cost_20g')
 elif n==8:chart(sl,C.sort_values('protein_per_100kcal',ascending=False),'protein_per_100kcal')
 elif n==9:
  for i,r in enumerate(C.itertuples()):
   y=3.1+i*.53;txt(sl,r.short,.7,y,1.8,.35,15,True);txt(sl,f'단백질 {r.protein_g}g',2.5,y,2.5,.35,15);txt(sl,f'{r.kcal}kcal',5.2,y,2.2,.35,15);txt(sl,'상세 라벨 미검증' if pd.isna(r.sugar_g) else f'당류 {r.sugar_g:g}g',8,y,3.3,.35,15)
  txt(sl,'셀렉스: WPI 사용 공식 안내 / 마이밀: 대두·완두·귀리 표기(재게시). 원료 유형은 우열 점수가 아니다.',.7,6.25,11.9,.5,11,False,MUTED)
 elif n in [10,11]:sl.shapes.add_picture(str(R/('analysis/figures/04_tradeoff.png' if n==10 else 'analysis/figures/06_lab2023.png')),Inches(.7),Inches(3.0),width=Inches(11.7),height=Inches(3.7))
 elif n==12:
  for i,(a,b,c) in enumerate([('COST FIRST','그린비아','관측 구성 중 20g당 약 1,052원'),('ENERGY FIRST','셀렉스','100kcal당 단백질 약 20.2g'),('BEFORE BUYING','표시·배송 재확인','현재 맛·수량·성분·총액 확인')]):card(sl,.6+i*4.2,3.3,3.98,a,b,c)
 else:
  for i,r in enumerate(S.itertuples()):
   x=.65+(i//6)*6.35;y=3+(i%6)*.51;sh=txt(sl,f'[{r.source_id}] {r.title[:39]}',x,y,5.95,.43,11,True);sh.text_frame.paragraphs[0].runs[0].hyperlink.address=r.url
  txt(sl,'가격·영양표는 스냅샷 / 원문 전문·인증 파일 비공개 / 2023 시험과 현재 가격 결합 금지',.7,6.4,11.9,.35,11,False,MUTED)
prs.save(P/'protein_value_analysis.pptx')
(P/'speaker_notes.md').write_text('\n\n'.join(f'## {i+1}. {s[1]}\n\n{s[3]}\n\n근거: {s[4]}' for i,s in enumerate(slides)),encoding='utf-8')
(P/'slides.json').write_text(json.dumps([dict(number=i+1,kicker=s[0],title=s[1],subtitle=s[2],notes=s[3],sources=s[4].split('|')) for i,s in enumerate(slides)],ensure_ascii=False,indent=2),encoding='utf-8')
data=json.loads((D/'data.json').read_text(encoding='utf-8'));data['slides']=json.loads((P/'slides.json').read_text(encoding='utf-8'))
template=(R/'scripts/template.html').read_text(encoding='utf-8');(D/'index.html').write_text(template.replace('__DATA__',json.dumps(data,ensure_ascii=False)),encoding='utf-8');(D/'.nojekyll').write_text('',encoding='utf-8')
src=R.parent/'coffee-brand-analysis/presentation/assets'
for name in ['anime.umd.min.js','anime-LICENSE.txt']:shutil.copy2(src/name,A/name)
print('Editable PPTX and animated 13-slide HTML generated')
